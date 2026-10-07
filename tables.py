"""Tables for people: every analysis result as rows of single-valued
cells, written as CSV files and as one workbook per experiment with a
Summary tab in front.

A Table is the typed form an analysis hands over; the human-readable
text and markdown a script also prints are renderings of the same
numbers. Cells hold one number or one string each, never a composite
like "168 (165-172)" or "15.8 / 17.1 / 17.7"; a statistic with a
spread is three columns. Column names carry the unit (`SourceF`,
`Gpm`, `HeatKbtuPerH`) in the sema camel form, so a CSV header reads
like a sema word's fields.

No sema word holds an analysis table; a word for tabular results
retires this module.
"""

import csv
from pathlib import Path
from typing import NamedTuple, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

Cell = int | float | str | None


class Table(NamedTuple):
    """One table of results: a sheet in the workbook, a file in csv/."""

    name: str  # sheet and file name, <= 31 characters, no / \\ ? * [ ] :
    title: str  # one line saying what the rows are
    columns: tuple[str, ...]
    rows: Sequence[tuple[Cell, ...]]
    source: str  # the script that produced it

    def check(self) -> "Table":
        assert len(self.name) <= 31 and not set(self.name) & set('/\\?*[]:'), self.name
        for r in self.rows:
            assert len(r) == len(self.columns), (self.name, r)
            for c in r:
                assert c is None or isinstance(c, (int, float, str)), (self.name, c)
        return self


def write_csv(table: Table, folder: Path) -> Path:
    folder.mkdir(exist_ok=True)
    path = folder / f"{table.name}.csv"
    with path.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(table.columns)
        for r in table.rows:
            w.writerow(["" if c is None else c for c in r])
    return path


def write_workbook(tables: Sequence[Table], path: Path, experiment: str, about: str) -> Path:
    """One workbook: a Summary tab naming every other tab, then one tab
    per table in the order given."""
    wb = Workbook()
    summary = wb.active
    assert summary is not None
    summary.title = "Summary"
    bold = Font(bold=True)
    summary.append([experiment])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([about])
    summary.append([])
    summary.append(["Tab", "What the rows are", "Rows", "Columns", "Produced by"])
    for c in range(1, 6):
        summary.cell(row=4, column=c).font = bold
    for t in tables:
        t.check()
        summary.append([t.name, t.title, len(t.rows), len(t.columns), t.source])
        ws = wb.create_sheet(t.name)
        ws.append(list(t.columns))
        for c in range(1, len(t.columns) + 1):
            ws.cell(row=1, column=c).font = bold
        for r in t.rows:
            ws.append(list(r))
        ws.freeze_panes = "A2"
        for i, col in enumerate(t.columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(col) + 2))
    for i, width in enumerate((34, 80, 8, 9, 28), start=1):
        summary.column_dimensions[get_column_letter(i)].width = width
    summary.freeze_panes = "A5"
    wb.save(path)
    return path
