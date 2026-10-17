"""Fit the star model (slab node, room branch, ground branch) to the
no-charge spans in the folder's `no.charge.*` gw.readings instances, then
check it by forward simulation and write `star-fit.xlsx`.

Stage one of `star-model.md`: one lump at the floor sensor,
    dT/dt = -a (T - T_room) - b (T - T_ground) + s S(t),
a shared across zones, b and T_ground per zone (T_ground per season), s
per zone with S the measured horizontal irradiance in W/m2 at the
Millinocket station (`irradiance-kmemilli18.csv.gz`, see pull_irradiance.py).
Fitted in integrated form at every one-minute sample,
    T(t) - T(0) = -a I1(t) - b I2(t) + (b T_ground) t + s I3(t),
with I1 = cumulative trapezoid of (T - T_room), I2 = of T, I3 = of S, so no
derivative of the data is ever taken. Ordinary least squares over every
span of every zone at once.

The check: each span simulated forward from its first sample with the
measured room temperature, exact discretisation at one minute, against
the measured floor. One sheet per span with a chart; a Summary sheet with
the parameters and per-span RMS.
"""

import datetime as dt
import gzip
import json
import math
import sys
from pathlib import Path
from typing import NamedTuple
from zoneinfo import ZoneInfo

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from pydantic import TypeAdapter

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parent / "src"))
sys.path.insert(0, str(HERE.parent))

from gwexp.sema.codec import SemaCodec  # noqa: E402
from gwexp.sema.property_format import LeftRightDot, SpaceheatName  # noqa: E402
from gwexp.sema.types import GwReadings  # noqa: E402
from naming import validate_lrd  # noqa: E402
from pull_irradiance import read_observations  # noqa: E402
from unit_encodings import to_display, word_encoding  # noqa: E402

SH = TypeAdapter(SpaceheatName)
LRD = TypeAdapter(LeftRightDot)
TA = validate_lrd("hw1.isone.me.versant.keene.spruce.ta")
ET = ZoneInfo("America/New_York")
DIST_FLOW = SH.validate_python("dist-flow")
FONT = "Arial"
BOLD = Font(name=FONT, bold=True)

# Instances, with the season each belongs to (ground temperature is per season).
INSTANCES: list[tuple[LeftRightDot, str]] = [
    (LRD.validate_python("no.charge.apr"), "spring"),
    (LRD.validate_python("no.charge.may"), "spring"),
    (LRD.validate_python("no.charge.fall"), "fall"),
]


class Zone(NamedTuple):
    """A slab zone: its label and the floor, room and heat-call channels."""

    label: str
    floor: SpaceheatName
    room: SpaceheatName
    call: SpaceheatName


ZONES = [
    Zone("bedrooms", SH.validate_python("zone1-bedrooms-floor-temp"),
         SH.validate_python("zone1-bedrooms-gw-temp"), SH.validate_python("zone1-bedrooms-heat-call")),
    Zone("living", SH.validate_python("zone2-living-rm-floor-temp"),
         SH.validate_python("zone2-living-rm-gw-temp"), SH.validate_python("zone2-living-rm-heat-call")),
    Zone("garage", SH.validate_python("zone4-garage-floor-temp"),
         SH.validate_python("zone4-garage-gw-temp"), SH.validate_python("zone4-garage-heat-call")),
]

# Room-sensor events (not physics) that break a span: see no-charge-spans.md.
EXCLUDE: dict[str, list[tuple[str, str]]] = {
    "bedrooms": [("2026-03-30 00:00", "2026-03-30 14:00"), ("2026-04-04 13:30", "2026-04-04 16:30")],
    "living": [("2026-03-30 00:00", "2026-03-30 14:00"), ("2026-05-05 11:00", "2026-05-07 12:00")],
    "garage": [("2026-03-30 00:00", "2026-03-30 14:00")],
}
MIN_SPAN_H = 2.0      # spans shorter than this are not used
BLIP_MIN = 60         # a charging blip shorter than this does not end a span
FLOW_ON_GPM = 0.05
STEP_MIN = 1


class Series(NamedTuple):
    """One channel's readings in natural units, time-ordered (ET, naive)."""

    times: list[dt.datetime]
    values: list[float]


class Span(NamedTuple):
    """A no-charge span on the one-minute grid: zone, season, times, floor and room."""

    zone: str
    season: str
    times: list[dt.datetime]
    floor: np.ndarray
    room: np.ndarray
    solar: np.ndarray  # horizontal irradiance W/m2

    @property
    def hours(self) -> float:
        return (self.times[-1] - self.times[0]).total_seconds() / 3600


class Fit(NamedTuple):
    """The fitted one-lump star: a shared, b and solar gain s per zone, T_ground per zone/season."""

    a: float
    b: dict[str, float]
    t_ground: dict[tuple[str, str], float]
    s: dict[str, float]  # F per (W h / m2) of horizontal irradiance


def load(codec: SemaCodec, condition: LeftRightDot) -> dict[SpaceheatName, Series]:
    path = HERE / f"{TA}-{condition}-gw.readings-000.json.gz"
    pull = codec.from_dict(json.loads(gzip.decompress(path.read_bytes())), expect=GwReadings)
    words = {c.name: c for c in pull.channels}
    out: dict[SpaceheatName, Series] = {}
    for cr in pull.channel_readings_list:
        encoding = word_encoding(words[cr.channel_name])
        pairs = sorted(zip(cr.scada_read_time_unix_ms_list, cr.value_list))
        times, values = [], []
        for t_ms, raw in pairs:
            v, _unit = to_display(raw, encoding)
            if cr.channel_name == DIST_FLOW and v > 15:  # the flow meter's nonsense spikes
                continue
            times.append(dt.datetime.fromtimestamp(t_ms / 1000, tz=ET).replace(tzinfo=None))
            values.append(v)
        out[cr.channel_name] = Series(times, values)
    return out


def grid_of(series: dict[SpaceheatName, Series]) -> list[dt.datetime]:
    t0 = min(s.times[0] for s in series.values()).replace(second=0, microsecond=0)
    t1 = max(s.times[-1] for s in series.values())
    n = int((t1 - t0).total_seconds() // 60) + 1
    return [t0 + dt.timedelta(minutes=i) for i in range(n)]


def on_grid(s: Series, grid: list[dt.datetime], hold_min: int, interpolate: bool) -> np.ndarray:
    """A series on the minute grid: linear interpolation (temperatures) or
    sample-and-hold (flow, calls), nan beyond hold_min from a reading."""
    g0 = grid[0]
    gx = np.arange(len(grid), dtype=float)
    sx = np.array([(t - g0).total_seconds() / 60 for t in s.times])
    sv = np.array(s.values)
    if interpolate:
        y = np.interp(gx, sx, sv, left=np.nan, right=np.nan)
    else:
        idx = np.searchsorted(sx, gx, side="right") - 1
        y = np.where(idx >= 0, sv[np.clip(idx, 0, len(sv) - 1)], np.nan)
    # distance to the nearest reading, in minutes
    left = np.searchsorted(sx, gx, side="right") - 1
    right = np.clip(left + 1, 0, len(sx) - 1)
    d_left = np.where(left >= 0, gx - sx[np.clip(left, 0, len(sx) - 1)], np.inf)
    d_right = np.abs(sx[right] - gx)
    y[np.minimum(d_left, d_right) > hold_min] = np.nan
    return y


def load_solar() -> Series:
    """The station's 5-minute irradiance highs, W/m2, ET naive like the readings."""
    times, values = [], []
    for o in read_observations():
        if o.solar_w_m2_high is None:
            continue
        times.append(dt.datetime.fromisoformat(o.obs_time_local))
        values.append(o.solar_w_m2_high)
    return Series(times, values)


def no_charge_spans(series: dict[SpaceheatName, Series], solar: Series, season: str) -> list[Span]:
    grid = grid_of(series)
    flow = on_grid(series[DIST_FLOW], grid, 10, False)
    sun = on_grid(solar, grid, 15, True)
    spans: list[Span] = []
    for z in ZONES:
        if z.floor not in series or z.room not in series:
            continue
        floor = on_grid(series[z.floor], grid, 10, True)
        room = on_grid(series[z.room], grid, 10, True)
        call = on_grid(series[z.call], grid, 10, False)
        charging = (flow > FLOW_ON_GPM) & (call > 0.5)
        known = ~np.isnan(flow) & ~np.isnan(call) & ~np.isnan(floor) & ~np.isnan(room) & ~np.isnan(sun)
        for a, b in EXCLUDE[z.label]:
            ta, tb = dt.datetime.fromisoformat(a), dt.datetime.fromisoformat(b)
            for i, t in enumerate(grid):
                if ta <= t < tb:
                    known[i] = False
        # drop charging blips shorter than BLIP_MIN
        i, n = 0, len(grid)
        while i < n:
            if charging[i]:
                j = i
                while j < n and charging[j]:
                    j += 1
                if j - i < BLIP_MIN:
                    charging[i:j] = False
                i = j
            else:
                i += 1
        usable = known & ~charging
        i = 0
        while i < n:
            if not usable[i]:
                i += 1
                continue
            j = i
            while j < n and usable[j]:
                j += 1
            if (j - i) * STEP_MIN / 60 >= MIN_SPAN_H:
                spans.append(Span(z.label, season, grid[i:j], floor[i:j].copy(), room[i:j].copy(), sun[i:j].copy()))
            i = j
    return spans


def cumtrapz(y: np.ndarray, dt_h: float) -> np.ndarray:
    out = np.zeros_like(y)
    out[1:] = np.cumsum((y[1:] + y[:-1]) / 2 * dt_h)
    return out


def fit(spans: list[Span], shared_a: bool = True) -> Fit:
    """OLS on the integrated model over all spans at once."""
    zones = sorted({s.zone for s in spans})
    seasons = sorted({(s.zone, s.season) for s in spans})
    a_cols = [None] if shared_a else zones
    nb, ntg = len(a_cols), len(a_cols) + len(zones)
    ns = ntg + len(seasons)
    ncol = ns + len(zones)
    rows, rhs = [], []
    dt_h = STEP_MIN / 60
    for s in spans:
        i1 = cumtrapz(s.floor - s.room, dt_h)
        i2 = cumtrapz(s.floor, dt_h)
        i3 = cumtrapz(s.solar, dt_h)
        t = np.arange(len(s.floor)) * dt_h
        x = np.zeros((len(s.floor), ncol))
        x[:, a_cols.index(None if shared_a else s.zone)] = -i1
        x[:, nb + zones.index(s.zone)] = -i2
        x[:, ntg + seasons.index((s.zone, s.season))] = t
        x[:, ns + zones.index(s.zone)] = i3
        rows.append(x)
        rhs.append(s.floor - s.floor[0])
    beta, *_ = np.linalg.lstsq(np.vstack(rows), np.concatenate(rhs), rcond=None)
    a = {z: beta[a_cols.index(None if shared_a else z)] for z in zones}
    b = {z: beta[nb + zones.index(z)] for z in zones}
    tg = {zs: beta[ntg + seasons.index(zs)] / b[zs[0]] for zs in seasons}
    sg = {z: beta[ns + zones.index(z)] for z in zones}
    return Fit(a[zones[0]] if shared_a else a, b, tg, sg)  # type: ignore[arg-type]


def simulate(s: Span, f: Fit) -> np.ndarray:
    a = f.a if isinstance(f.a, float) else f.a[s.zone]
    b = f.b[s.zone]
    tg = f.t_ground[(s.zone, s.season)]
    k = a + b
    decay = math.exp(-k * STEP_MIN / 60)
    y = np.empty_like(s.floor)
    y[0] = s.floor[0]
    sg = f.s[s.zone]
    for i in range(1, len(y)):
        teq = (a * s.room[i - 1] + b * tg + sg * s.solar[i - 1]) / k
        y[i] = teq + (y[i - 1] - teq) * decay
    return y


def rms(x: np.ndarray) -> float:
    return float(np.sqrt(np.mean(x * x)))


def write_workbook(spans: list[Span], f: Fit, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["One-lump star fitted on no-charge spans by cumulative-integral least squares; see star-model.md"])
    ws.append([])
    ws.append(["parameter", "zone", "season", "value", "unit"])
    ws.append(["a = 1/(R_room C)", "shared", "", round(float(f.a), 5), "per hour"])
    for z, b in f.b.items():
        ws.append(["b = 1/(R_ground C)", z, "", round(float(b), 5), "per hour"])
    for (z, se), tg in f.t_ground.items():
        ws.append(["T_ground", z, se, round(float(tg), 1), "F"])
    for z, sg in f.s.items():
        ws.append(["s = solar gain", z, "", round(float(sg) * 1000, 4), "F per kWh/m2 horizontal"])
    ws.append(["time constant 1/(a+b)", "", "", "", "hours"])
    for z, b in f.b.items():
        ws.append(["", z, "", round(1 / (float(f.a) + float(b)), 1), "hours"])
    ws.append([])
    ws.append(["span sheet", "zone", "season", "start (ET)", "hours", "floor start", "room start", "RMS model-actual F", "max abs F"])
    for cell in ws[3] + ws[ws.max_row]:
        cell.font = BOLD
    all_res = []
    for s in spans:
        y = simulate(s, f)
        res = y - s.floor
        all_res.append(res)
        name = f"{s.zone} {s.times[0]:%m-%d %H%M}"
        ws.append([name, s.zone, s.season, s.times[0], round(s.hours, 1), round(float(s.floor[0]), 1),
                   round(float(s.room[0]), 1), round(rms(res), 2), round(float(np.max(np.abs(res))), 2)])
        sh = wb.create_sheet(name[:31])
        sh.append(["time (ET)", "room F", "floor actual F", "floor model F", "model - actual F", "solar W/m2"])
        for cell in sh[1]:
            cell.font = BOLD
        step = 5  # rows every 5 minutes keep the sheets light; the fit used every minute
        for i in range(0, len(s.times), step):
            sh.append([s.times[i], round(float(s.room[i]), 2), round(float(s.floor[i]), 2),
                       round(float(y[i]), 2), round(float(res[i]), 2), round(float(s.solar[i]), 0)])
        sh.column_dimensions["A"].width = 18
        for r in sh.iter_rows(min_row=2, max_col=1):
            r[0].number_format = "mm-dd hh:mm"
        ch = LineChart()
        ch.title = f"{s.zone} from {s.times[0]:%Y-%m-%d %H:%M}: floor actual vs star model"
        ch.y_axis.title = "F"
        ch.x_axis.title = "time"
        ch.height, ch.width = 10, 26
        ch.add_data(Reference(sh, min_col=2, max_col=4, min_row=1, max_row=sh.max_row), titles_from_data=True)
        ch.set_categories(Reference(sh, min_col=1, min_row=2, max_row=sh.max_row))
        ch.x_axis.number_format = "mm-dd hh"
        for ser in ch.series:
            ser.smooth = False
            ser.graphicalProperties.line.width = 15000
        sh.add_chart(ch, "G2")
    ws.append([])
    ws.append(["all spans", "", "", "", "", "", "", round(rms(np.concatenate(all_res)), 2)])
    for col in range(1, 10):
        ws.column_dimensions[get_column_letter(col)].width = 20
    wb.save(path)


def main() -> None:
    codec = SemaCodec()
    solar = load_solar()
    spans: list[Span] = []
    for condition, season in INSTANCES:
        spans += no_charge_spans(load(codec, condition), solar, season)
    spans.sort(key=lambda s: (s.zone, s.times[0]))
    f = fit(spans)
    total = sum(s.hours for s in spans)
    print(f"{len(spans)} spans, {total:.0f} h; a={f.a:.4f}/h; b=" +
          ", ".join(f"{z} {b:.4f}" for z, b in f.b.items()) + "; T_ground=" +
          ", ".join(f"{z}/{se} {tg:.1f}F" for (z, se), tg in f.t_ground.items()) +
          "; s (F per kWh/m2)=" + ", ".join(f"{z} {sg * 1000:.3f}" for z, sg in f.s.items()))
    out = HERE / "star-fit.xlsx"
    write_workbook(spans, f, out)
    res = np.concatenate([simulate(s, f) - s.floor for s in spans])
    print(f"forward-simulation RMS over all spans {rms(res):.2f} F; wrote {out}")
    f2 = fit(spans, shared_a=False)
    res2 = np.concatenate([simulate(s, f2) - s.floor for s in spans])
    print("per-zone a: " + ", ".join(f"{z} {a:.4f}" for z, a in f2.a.items()) + f"; RMS {rms(res2):.2f} F")  # type: ignore[union-attr]


if __name__ == "__main__":
    main()
